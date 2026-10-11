#!/usr/bin/env python3
"""Create editable DOCX and PDF application drafts from explicitly verified facts.

No claims are invented. Missing resume facts remain clearly marked for review.
Requires python-docx and reportlab for the corresponding output formats.
"""
import argparse
import json
from pathlib import Path
from datetime import date

BASE = Path.home() / "JobSearchAgent"
PROFILE = BASE / "candidate_profile.json"
OUT = BASE / "application_drafts"

def load(path, default):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default

def build_text(job, facts):
    title = job.get("title") or "[VERIFY JOB TITLE]"
    company = job.get("company") or "[VERIFY COMPANY]"
    lines = [
        "APPLICATION DRAFT — REVIEW BEFORE USE",
        f"Role: {title}",
        f"Company: {company}",
        f"Prepared: {date.today().isoformat()}",
        "",
        "COVER LETTER DRAFT",
        f"Dear Hiring Team at {company},",
        "",
        f"I am interested in the {title} opportunity. The points below are drawn only from the verified facts supplied to JobSearchAgent; please confirm they accurately represent your experience before sending.",
        ""
    ]
    verified = facts.get("verified_facts", [])
    if verified:
        lines.extend(["Relevant verified experience:"])
        lines.extend([f"- {str(x)}" for x in verified])
    else:
        lines.extend([
            "[ADD verified experience from your current resume before sending.]",
            "[ADD one specific example that matches the job requirements.]"
        ])
    lines.extend([
        "",
        "[ADD a concise, truthful paragraph explaining why this role and company fit.]",
        "",
        "Thank you for your consideration.",
        "[YOUR NAME]",
        "[YOUR CONTACT DETAILS]",
        "",
        "TAILORED RESUME CHECKLIST",
        "- Reorder existing resume bullets to emphasize relevant verified evidence.",
        "- Add measurable outcomes only where supported by the original resume.",
        "- Do not add skills, credentials, employers, dates, or achievements that cannot be verified.",
        "",
        "JOB DETAILS TO VERIFY",
        f"Posting URL: {job.get('url') or '[ADD VERIFIED JOB URL]'}",
        f"Location: {job.get('location') or '[VERIFY LOCATION]'}",
        f"Posting date: {job.get('posted_date') or '[UNKNOWN — VERIFY ON ORIGINAL LISTING]'}",
        "",
        "DRAFT ONLY — this file has not been submitted."
    ])
    return lines

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--job-json", required=True, help="JSON file containing title, company, location, posted_date, url")
    parser.add_argument("--profile-json", default=str(PROFILE), help="Profile JSON; use verified_facts for resume-grounded statements")
    parser.add_argument("--name", default="application-draft", help="Output filename stem")
    args = parser.parse_args()
    job = load(args.job_json, {})
    facts = load(args.profile_json, {})
    lines = build_text(job, facts)
    OUT.mkdir(parents=True, exist_ok=True)
    stem = OUT / Path(args.name).name
    textfile = stem.with_suffix(".txt")
    textfile.write_text("\n".join(lines) + "\n", encoding="utf-8")
    created = [str(textfile)]
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("Application Draft — Review Before Use", 0)
        for line in lines[1:]:
            if line in ("COVER LETTER DRAFT", "TAILORED RESUME CHECKLIST", "JOB DETAILS TO VERIFY"):
                doc.add_heading(line.title(), level=1)
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line:
                doc.add_paragraph(line)
            else:
                doc.add_paragraph("")
        p = stem.with_suffix(".docx")
        doc.save(p)
        created.append(str(p))
    except ImportError:
        print("DOCX skipped: install python-docx to enable it.")
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from xml.sax.saxutils import escape
        p = stem.with_suffix(".pdf")
        styles = getSampleStyleSheet()
        story = []
        for line in lines:
            if line in ("COVER LETTER DRAFT", "TAILORED RESUME CHECKLIST", "JOB DETAILS TO VERIFY"):
                story.append(Paragraph(escape(line.title()), styles["Heading2"]))
            elif line:
                story.append(Paragraph(escape(line), styles["BodyText"]))
            story.append(Spacer(1, 0.08 * inch))
        SimpleDocTemplate(str(p), pagesize=letter).build(story)
        created.append(str(p))
    except ImportError:
        print("PDF skipped: install reportlab to enable it.")
    print("Created review-only draft files:")
    print("\n".join(created))
    print("Verify every statement and replace all placeholders before use.")

if __name__ == "__main__":
    main()
